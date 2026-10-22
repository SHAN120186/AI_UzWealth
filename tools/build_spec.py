from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from pathlib import Path
import json

ROOT=Path('/workspace/scratch/2707cec583ab')
OUT=ROOT/'output/UzWealth_AI_Technical_Specification_v1.docx'
W=ROOT/'work'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    styles[name].font.name='Calibri';styles[name].font.color.rgb=RGBColor(0,0,0)
styles['Normal'].font.size=Pt(11)
styles['Normal'].paragraph_format.space_after=Pt(6)
styles['Normal'].paragraph_format.line_spacing=1.06
styles['Title'].font.size=Pt(27);styles['Title'].font.bold=True
styles['Title'].paragraph_format.space_after=Pt(12)
styles['Subtitle'].font.size=Pt(13)
styles['Subtitle'].font.italic=False
# Remove paragraph borders inherited from the default Word template.
for sty in styles:
    for border in list(sty.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name,size,bef,aft in [('Heading 1',17,0,10),('Heading 2',12.5,11,6),('Heading 3',11,8,4)]:
    s=styles[name];s.font.size=Pt(size);s.font.bold=True
    s.paragraph_format.space_before=Pt(bef);s.paragraph_format.space_after=Pt(aft)
    s.paragraph_format.keep_with_next=True
code_style=styles.add_style('Code',1)
code_style.font.name='DejaVu Sans Mono';code_style.font.size=Pt(8.3)
code_style.paragraph_format.space_after=Pt(0)
code_style.paragraph_format.line_spacing=1.0
code_style.paragraph_format.keep_with_next=False
small=styles.add_style('Small',1);small.font.name='Calibri';small.font.size=Pt(9)
small.paragraph_format.space_after=Pt(4)
header=sec.header.paragraphs[0]
header.text='UzWealth AI   |   Техническое задание   |   Версия 1.0'
header.style='Small'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.style='Small';footer.add_run('Страница ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title='Техническое задание UzWealth AI'
doc.core_properties.subject='Нормативная спецификация MVP и критерии приемки'
doc.core_properties.author='';doc.core_properties.keywords='UzWealth AI, LangGraph, FastAPI, Telegram Mini App'

def p(t,style=None): return doc.add_paragraph(t,style)
def h(t,level=2):return doc.add_heading(t,level)
def page(t):
    par=h(t,1);par.paragraph_format.page_break_before=True
def bullets(items):
    for x in items:p(x,'List Bullet')
def nums(items):
    for i,x in enumerate(items,1):p(f'{i}. {x}')
def code(s):
    for line in s.strip('\n').split('\n'):p(line,'Code')
    p('','Small')
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,x in zip(t.rows[0].cells,headers):c.text=x
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for vals in rows:
        cells=t.add_row().cells
        for c,x in zip(cells,vals):c.text=str(x)
    for ri,row in enumerate(t.rows):
        rpr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');rpr.append(no)
        if ri==0:
            rep=OxmlElement('w:tblHeader');rpr.append(rep)
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for ed in ['top','bottom','left','right']:
                z=OxmlElement('w:'+ed);z.set(qn('w:w'),'90');z.set(qn('w:type'),'dxa');mar.append(z)
            tcpr.append(mar)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243746' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.02
                if headers[ci]=='Стоимость UZS':par.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                elif 'bps' in headers[ci] or headers[ci]=='ID':par.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in par.runs:
                    r.font.size=Pt(9.5);r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
                    r.font.bold=ri==0
    p('','Small')
    return t

def flow_image():
    im=Image.new('RGB',(1600,1380),'white');d=ImageDraw.Draw(im)
    font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',25)
    bold=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',27)
    def box(x,y,w,hh,lines,fill='#F2F5F7'):
        d.rounded_rectangle((x,y,x+w,y+hh),radius=12,fill=fill,outline='#73818B',width=2)
        for i,line in enumerate(lines):
            f=bold if i==0 else font
            bb=d.textbbox((0,0),line,font=f);d.text((x+(w-(bb[2]-bb[0]))/2,y+17+i*36),line,font=f,fill='#111111')
    def arr(points):
        d.line(points,fill='#485E6B',width=4)
        x,y=points[-1];px,py=points[-2]
        if y>py:d.polygon([(x,y),(x-9,y-15),(x+9,y-15)],fill='#485E6B')
        elif x>px:d.polygon([(x,y),(x-15,y-9),(x-15,y+9)],fill='#485E6B')
        else:d.polygon([(x,y),(x+15,y-9),(x+15,y+9)],fill='#485E6B')
    box(430,20,650,100,['START','Авторизация и схема запроса'])
    box(430,180,650,110,['1 extract_profile','Форма или LLM и подтверждение'])
    box(20,185,330,105,['needs_input','Уточнение и checkpoint'])
    box(430,360,650,110,['2 get_market_data','Redis или PostgreSQL'])
    box(1160,365,410,110,['failed','Нет обязательных данных'])
    box(430,540,650,110,['3 quant_allocation','Расчёт и проверка инвариантов'])
    box(1160,545,410,110,['failed','Баланс или лимиты нарушены'])
    box(430,720,650,110,['4 generate_report','Пояснение или статический текст'])
    box(430,920,650,110,['Публикация результата','Шаблон чисел и дисклеймера'])
    box(430,1110,650,110,['END','completed или with_warnings'])
    for a,b in [(120,180),(290,360),(470,540),(650,720),(830,920),(1030,1110)]:arr([(755,a),(755,b)])
    arr([(430,230),(350,230)])
    arr([(350,260),(390,260),(390,310),(755,310),(755,360)])
    arr([(1080,410),(1160,410)]);arr([(1080,595),(1160,595)])
    d.text((35,1270),'На каждой границе узла: ownership, deadline, cancel и runtime validation',font=font,fill='#333333')
    im.save(W/'graph.png')
flow_image()

h('Техническое задание UzWealth AI',0)
p('Telegram Mini App для расчёта сценариев размещения капитала в Узбекистане','Subtitle')
p('Версия 1.0   •   22 октября 2026 года   •   Объём поставки MVP','Small')
p('Разработать сервис, который принимает подтверждённый профиль пользователя, выбирает допустимые инструменты из проверенного каталога и детерминированно рассчитывает суммы, количества и денежный остаток. LLM извлекает параметры из текста и объясняет уже проверенный результат. Все денежные операции, ограничения и юридические тексты выполняются кодом.')
h('Статус и применение')
p('Документ задаёт единую реализацию для команды backend, frontend, QA и эксплуатации. Требования со словами «должен» и «запрещено» обязательны для MVP. Числовые лимиты и матрица риска ниже являются проектной базовой политикой для разработки и тестирования. Публичный запуск разрешается после согласований раздела 22; эти согласования не блокируют разработку на тестовых данных.')
table(['Решение','Зафиксированная реализация'],[
('Оркестрация','LangGraph с четырьмя бизнес узлами и условными переходами'),
('Расчёты','Чистый Python, int в тийинах и Decimal; матрица правил с версией'),
('Клиент','Telegram Mini App на React и TypeScript; русский и узбекский латиницей'),
('Сервер','FastAPI, PostgreSQL с расширением TimescaleDB, Redis, Celery'),
('Хранение состояния','PostgresSaver и предметные таблицы PostgreSQL'),
('Результат','Расчётный сценарий без открытия счетов, платежей и размещения заявок')],[1.5,5.5])
h('Граница поставки')
p('MVP включает сумовые депозиты, физические слитки золота и опциональный BTC. Криптовалюта по умолчанию выключена до допуска провайдера и правового согласования. Реальные пользовательские балансы, подключение брокерских счетов, автоторговля, ETH, альткоины, прогноз цены и оптимизация Марковица в поставку не входят. Векторная база и новостной RAG не требуются.')
p('Навигация: разделы 2–4 описывают поведение и архитектуру; 5–10 — профиль, Quant Core и рынок; 11–15 — контракты и хранение; 16–19 — безопасность и эксплуатацию; 20–23 — приёмку, поставку и источники.','Small')

page('2 Пользовательские сценарии')
h('Основной сценарий')
nums(['Пользователь открывает Mini App. Backend проверяет Telegram initData и создаёт сессию. Пользователь выбирает язык и принимает действующие версии условий и политики обработки данных.',
'Форма принимает выделенный инвестиционный бюджет, горизонт, ответы анкеты, допустимую потерю и ограничения. Пользователь подтверждает, что резерв на повседневные расходы уже выделен отдельно.',
'Кнопка «Рассчитать сценарий» создаёт фоновую задачу. Экран показывает проверку профиля, чтение рыночного среза, расчёт и подготовку объяснения.',
'Результат показывает целевые доли, рассчитанные количества, комиссии, остаток, источники и время котировок. Пользователь может сравнить новый сценарий с предыдущим, создав новый расчёт.'])
h('Альтернативный ввод')
p('Текст «100 млн сум на год с умеренным риском» является черновиком профиля. LLM возвращает только найденные значения и вопросы. Незаполненные поля остаются null. Форма подтверждения обязательна; до подтверждения финансовый расчёт не начинается. Фраза без валюты, конфликт сумм или риск без анкеты переводят задачу в needs_input.')
table(['Код','Обязательное поведение'],[
('FR 01','Редактирование суммы, горизонта или ограничения создаёт новый job; старый результат сохраняется.'),
('FR 02','Ошибка генерации текста сохраняет проверенную таблицу и включает стандартное объяснение.'),
('FR 03','Повторное нажатие с тем же ключом возвращает тот же job без нового расчёта.'),
('FR 04','Недоступность источника показывает возраст среза или отказ; скрытая подстановка демонстрационных цен запрещена.'),
('FR 05','Пользователь видит причины исключения инструмента и остаток денег. Наличие у банка или площадки требует отдельной проверки.'),
('FR 06','Все числовые поля и дисклеймер формирует серверный шаблон; перевод не меняет значения.')],[.8,6.2])
h('Требования к интерфейсу')
p('Поддержать экран шириной от 360 px, состояния loading, needs_input, completed, failed и cancelled. Кнопки имеют текстовые подписи, поля — явные единицы. Ошибка показывается рядом с полем. Цвет не является единственным способом обозначения риска. Дата и время в интерфейсе — Asia/Tashkent; в API — UTC. Диаграмма долей включает cash и поясняет, что явные комиссии показаны отдельно.')

page('3 Компоненты и границы ответственности')
p('Развёртывание MVP — модульный backend в контейнерах. FastAPI принимает запросы, а Celery выполняет граф вне HTTP обработчика. Сбор данных запускается отдельным расписанием. При расчёте пользователя запрещены парсинг сайтов и обращение к внешним рыночным API.')
table(['Компонент','Ответственность','Сохранность'],[
('Mini App','Форма, подтверждение профиля, статусы, результат','Не хранит секреты и финансовую историю локально'),
('FastAPI','Auth, валидация, ownership, job API и SSE','Job и outbox фиксируются до ответа 202'),
('Worker LangGraph','Четыре узла, проверки и подготовка результата','Checkpoint после каждого завершённого узла'),
('Quant Core','Бюджеты, лоты, комиссии и инварианты','Версия алгоритма и входов сохраняется'),
('Market worker','Плановый сбор, нормализация, контроль качества','Сначала БД, затем публикация в Redis'),
('PostgreSQL','Профили, задания, результаты, каталог, события','Единственный источник истины приложения'),
('TimescaleDB','История котировок в том же PostgreSQL','Hypertable только для временных рядов'),
('Redis','Кэш, rate limits, Celery broker, уведомления SSE','Потеря Redis не удаляет job или результат'),
('LLM adapter','Извлечение текста и пояснения по схеме','Не имеет инструментов торговли и записи финансовых полей')],[1.25,3.0,2.75])
h('Потоки данных')
p('Mini App обращается только к FastAPI по HTTPS. API записывает job и outbox одной транзакцией. Dispatcher доставляет job в Celery. Worker читает закреплённые правила и каталог, выполняет LangGraph, записывает результат и событие завершения. SSE читает сохранённые события; Redis ускоряет уведомление о новых событиях.')
p('Market worker получает данные из разрешённых источников, сохраняет сырые ответы и нормализованные записи, создаёт неизменяемый snapshot, а затем атомарно обновляет указатель latest в Redis. Частичный или невалидный сбор не перезаписывает последний пригодный срез.')
h('Разделение сред')
p('Development, test и production используют отдельные базы, bot token и ключи LLM. Production не читает тестовые fixtures. Данные пользователей и секреты не включаются в Docker образ, журнал сборки, checkpoint или трассировку модели.')

page('4 Граф выполнения и обработка ошибок')
par=p('');par.add_run().add_picture(str(W/'graph.png'),width=Inches(6.8))
p('Рисунок 1. Четыре бизнес узла LangGraph и основные условные переходы.','Small')
p('extract_profile пропускает LLM для формы. В текстовом режиме узел сохраняет черновик и останавливается на уточнении. После подтверждения продолжает тот же job с новой input_version. Для нового сценария создаётся новый job и новый snapshot.')
p('get_market_data допускает старый срез только в пределах жёсткого возраста раздела 9. quant_allocation передаёт управление generate_report только после успешной проверки всех инвариантов. generate_report при ошибке LLM формирует статический отчёт и completed_with_warnings. Невалидный финансовый результат никогда не публикуется.')

page('5 Подтверждённый профиль и оценка риска')
table(['Поле ConfirmedProfileV1','Тип и правило'],[
('capital_tiyin','Строка целого числа; 100 000 000–100 000 000 000 000 тийинов, то есть 1 млн–1 трлн UZS'),
('horizon_months','Строго int, 1–60; срок задаётся пользователем'),
('selected_risk','conservative | moderate | aggressive'),
('loss_capacity','0 — потеря нарушит обязательства; 1 — допустима часть; 2 — бюджет допускает существенную потерю'),
('loss_behavior','0 — вывести при падении; 1 — удержать без докупки; 2 — допускает удержание при глубоком падении'),
('experience','0 — нет опыта; 1 — депозиты и золото; 2 — понимает волатильные активы и комиссии'),
('max_loss_bps','int 0–10000; допустимая пользователем потеря; не гарантируемый предел просадки'),
('crypto_opt_in','bool; отдельное согласие на рассмотрение BTC'),
('resident_uz и adult_confirmed','bool; MVP доступен только заявившим резидентство РУз и совершеннолетие'),
('reserve_confirmed','true; бюджет уже отделён от жизненного резерва'),
('risk_acknowledged','Строго true в подтверждённом профиле; явное принятие возможности убытка'),
('confirmed_at и profile_version','UTC datetime и int; выставляет backend после подтверждения')],[2.1,4.9])
h('Правила класса риска')
p('Сумма трёх ответов 0–2 означает conservative, 3–4 — moderate, 5–6 — aggressive. effective_risk — минимум по порядку conservative < moderate < aggressive из класса анкеты, selected_risk, лимита потери и горизонта. max_loss_bps меньше 1000 ограничивает класс до conservative; 1000–2499 — до moderate. Для горизонта меньше 24 месяцев класс не выше moderate; короткий срок дополнительно ограничивает активы. Пороги подлежат утверждению владельцем риска.')
p('Если loss_capacity равен 0, max_loss_bps равен 0, reserve_confirmed не true либо пользователь не подтверждает понимание возможности убытка, система возвращает needs_input с причиной SUITABILITY_NOT_CONFIRMED. Она не предлагает портфель с обещанием отсутствия потерь. Противоречивые ответы показываются пользователю для исправления.')
h('Допуск криптовалюты')
p('BTC рассматривается только при crypto_opt_in=true, experience=2, горизонте от 12 месяцев, max_loss_bps от 2000 и допущенном провайдере. Отсутствие любого условия даёт криптодолe 0. Telegram подтверждает происхождение сеанса, но не заменяет KYC провайдера. Пользовательский текст не может снять ограничение.')

page('6 Версионируемая политика распределения')
p('Для разработки зафиксировать risk_policy_v1_demo. Это матрица правил, а не оптимизация Марковица и не оценка ожидаемой доходности. Переход к статистическому оптимизатору является отдельным изменением методологии с историческими данными, проверкой устойчивости и новой версией политики.')
table(['Класс','Депозит bps','Золото bps','BTC bps','Cash bps'],[
('conservative',8000,2000,0,0),('moderate',6000,2500,1500,0),('aggressive',5000,2500,2500,0)],[1.6,1.35,1.35,1.35,1.35])
p('1 bps = 0,01 процентного пункта. Сумма целевых долей после применения ограничений равна 10000 bps. Порядок правил фиксирован: класс риска → горизонт → допуск активов → доступность данных → доступность продукта. Любая снятая доля переходит в cash, остальные доли не нормализуются вверх.')
table(['Условие','Ограничение'],[
('Горизонт 1–5 месяцев','Золото 0, BTC 0; освободившаяся доля в cash'),
('Горизонт 6–11 месяцев','Золото максимум 1000 bps, BTC 0'),
('Горизонт от 12 месяцев','Лимиты строки effective_risk и индивидуальный допуск'),
('BTC не допущен','Вся целевая доля BTC переходит в cash с кодом причины'),
('Нужные данные отсутствуют','Расчёт блокируется; автоматическая подмена инструмента запрещена'),
('Нет подходящего продукта','Бюджет класса остаётся cash; статус completed_with_warnings'),
('Все классы недоступны','ALLOCATION_INFEASIBLE; отдельного сценария «100% cash» как инвестиционного предложения нет')],[2.1,4.9])
h('Точность и сохранение бюджета')
p('Деньги хранятся как целое число тийинов. JSON передаёт их строкой, чтобы JavaScript не терял точность. Цены, ставки и количества передаются десятичными строками и обрабатываются Decimal. float, NaN, Infinity и отрицательные суммы запрещены. Все комиссии включаются в бюджет класса.')
code('''# Для deposit_uzs, gold_cbu, crypto_btc:
budget[i] = (capital_tiyin * target_share_bps[i]) // 10000
cash_budget = capital_tiyin - sum(budget)
# Любые дробные тийины распределения остаются в cash.
(cost[i] + fee[i]) * 10000 <= cap_bps[i] * capital_tiyin''')
p('Hard caps относятся к полной стоимости класса с комиссиями относительно исходного капитала C и проверяются точным целочисленным неравенством выше. cap_bps равен доле после всех ограничений. Net cost shares предназначены только для отображения и имеют другую базу. asset_id: cash_uzs, crypto_btc, deposit_uzs, gold_cbu. После расчёта лотов весь остаток в cash.')

page('7 Количества продуктов и комиссии')
h('Депозиты')
p('Каталог должен содержать валюту, минимальную сумму, шаг пополнения, сроки в месяцах, тип ставки, порядок начисления и выплаты, требования к клиенту, комиссии и условия досрочного закрытия. В MVP доступны только UZS продукты с фиксированной ставкой и сроком не больше горизонта. При неизвестных условиях продукт исключается.')
p('MVP допускает один режим: простые проценты без капитализации с выплатой в конце срока. Выбрать срок, наиболее близкий к горизонту снизу, затем большую номинальную ставку, затем product_id по возрастанию. Сумма депозита — максимальное целое число шагов, при котором principal + fee ≤ бюджет. Один продукт на job; дробления по банкам нет. Ниже минимальной суммы — cash.')
h('Физическое золото')
p('Цены задаются за целый слиток конкретного веса, а не за грамм. Использовать ask для покупки и отдельные bid для оценки выкупа. Допускаются только номиналы из проверенного каталога, соответствующие опубликованным продуктам ЦБ. Наличие unknown допускает лишь условный сценарий с предупреждением AVAILABILITY_UNCONFIRMED; sold_out исключается.')
p('Из всех допустимых целочисленных комбинаций в бюджете выбрать максимальный общий вес; затем минимальную полную стоимость с комиссиями; затем меньшее число слитков; затем лексикографически меньший отсортированный список product_id. Учитывать известные stock limits. Реализация — точный ограниченный поиск с отсечениями; не жадный алгоритм. Превышение лимита расчёта 2 секунды даёт QUANT_TIMEOUT, без приблизительной покупки.')
h('BTC')
p('MVP использует один спотовый UZS/BTC продукт допущенного провайдера. Плечо, деривативы и автопокупка запрещены. Нет поддерживаемого продукта в утверждённом каталоге — бюджет в cash. Для уже допущенного продукта отсутствие или просрочка ask либо тарифа даёт MARKET_DATA_UNAVAILABLE. min_notional и quantity_step обязательны в каталоге. Мировая цена BTC не заменяет локальный ask.')
code('''cost(q) = ceil_to_tiyin(q * ask_uzs_per_btc * 100)
fee(q) = ceil_to_tiyin(cost(q) * fee_rate + fixed_fee_tiyin)
q = n * quantity_step, n >= 0
# Найти максимальное целое n: cost(q) + fee(q) <= budget.
# Если cost(q) < min_notional: q = 0 и fee = 0.''')
p('Комиссия должна быть монотонной для бинарного поиска количества. Другой тариф требует отдельного адаптера и тестов. В тариф включаются применимые расходы выбранного пути покупки; неизвестная комиссия не принимается равной нулю. Вывод и перевод между кошельками в MVP не моделируются и явно исключены из расчёта.')

page('8 Проверочный расчёт на 100 млн сумов')
p('Фикстура F100 служит только для автоматизированной приёмки. Цены и комиссия ниже не являются котировками на дату документа. Профиль: 100 млн UZS, 12 месяцев, итоговый moderate, max_loss_bps=2000, experience=2, все подтверждения даны. BTC разрешён тестовой конфигурацией; целевые доли 6000 / 2500 / 1500 bps.')
table(['Тестовое условие','Значение'],[
('Депозит','dep_f100; минимум 1 млн UZS; шаг 0,01 UZS; 12 месяцев; 22% простых годовых с выплатой в конце; комиссия 0'),
('Золото ask','5 г = 5 750 000; 10 г = 11 500 000; 20 г = 23 000 000 UZS'),
('Золото bid','Целая упаковка 5/10/20 г: 5,25/10,5/21 млн UZS; повреждённая: 5/10/20 млн UZS'),
('BTC ask и шаг','1 000 000 000 UZS за BTC; шаг 0.00000001 BTC; min_notional 100 000 UZS'),
('Комиссия BTC','1% стоимости; фиксированная комиссия 0; оплата в UZS')],[1.65,5.35])
table(['Позиция','Количество','Стоимость UZS'],[
('Депозит','Один продукт','60 000 000,00'),('Золото','Один слиток 20 г','23 000 000,00'),
('BTC','0.01485148 BTC','14 851 480,00'),('Явная комиссия BTC','1%','148 514,80'),
('Денежный остаток','Золото 2 млн и BTC 5,20','2 000 005,20'),('Итого','Бюджет сохранён','100 000 000,00')],[2.0,2.25,2.75])
p('gold_5, gold_10, gold_20: комиссия 0, availability=available, по 100 штук. BTC product_id=btc_f100; фиксированный test_clock и все времена котировок 2026-01-15T06:00:00Z. Следующий шаг 0.01485149 BTC с комиссией стоит 15 000 004,90 UZS и превышает бюджет класса.')
h('Доли и оценка выкупа')
p('Доли по себестоимости рассчитываются от 99 851 485,20 UZS, то есть капитала за вычетом явных комиссий. После метода наибольших остатков: депозит 6009, золото 2304, BTC 1487, cash 200 bps. Сумма ровно 10000. Это не рыночные веса и не гарантия риск экспозиции. На экране отдельно показываются исходные целевые доли и явная комиссия.')
p('Оценка немедленного выкупа золота в фикстуре равна 21 млн UZS. Разница с ценой покупки 2 млн UZS отражается отдельно как спред, а не как повторная комиссия в бюджетном балансе.')

page('9 Контракт рыночного среза и свежесть')
p('Каждый job использует один MarketSnapshotV1. Snapshot неизменяемый и содержит ссылки на конкретные версии котировок, каталога и комиссий. При повторном выполнении job новые цены не подмешиваются. Время загрузки не заменяет время наблюдения или дату действия значения.')
table(['Объект','Обязательные поля'],[
('Snapshot','snapshot_id UUID; created_at UTC; catalog_version; provider_registry_version; quotes[]; schema_version=1.0'),
('Quote','quote_id; source_id; source_url; instrument_id; kind; value decimal string; currency; unit; observed_at nullable; effective_date nullable; fetched_at; verified_at; quality'),
('quality','valid | rejected; rejected записи не участвуют в snapshot для расчёта'),
('Время значения','observed_at либо effective_date обязателен; для ручного банковского продукта обязателен verified_at и дата условий'),
('kind и unit','fx_reference UZS/unit; gold_ask и gold_bid UZS/bar; crypto_ask UZS/BTC; deposit_nominal_rate fraction/year; inflation_yoy fraction/year')],[1.5,5.5])
table(['Набор','Опрос','Мягкий возраст','Жёсткая граница'],[
('Курс ЦБ','Каждый час','Следующая ожидаемая публикация','Более 3 рабочих дней задержки публикации'),
('Слитки ЦБ','Каждые 15 минут','2 часа после проверки','Более 1 рабочего дня без успешной проверки'),
('BTC ask и тариф','Каждые 30 секунд','60 секунд для ask','300 секунд для ask; тариф 24 часа'),
('Условия вкладов','Ежесуточно','24 часа','72 часа после ручной или автоматической проверки'),
('Инфляция','Ежесуточно','До следующей месячной публикации','45 дней после даты последней публикации'),
('Реестр провайдеров','Ежесуточно','24 часа','24 часа без подтверждения статуса')],[1.25,1.2,2.15,2.4])
p('Это проектные правила допустимости, а не SLA источников. Производственный календарь и ожидаемое время публикаций версионируются. На выходном последняя действующая запись ЦБ допустима до следующего ожидаемого обновления с указанием даты. Повторное скачивание той же страницы не делает старый ask BTC свежим.')
p('Возраст между мягкой и жёсткой границей даёт STALE_DATA_USED. Превышение жёсткой границы обязательного набора даёт MARKET_DATA_UNAVAILABLE. Обязательны только данные для ненулевых долей после eligibility. Инфляция и справочный FX необязательны, если не участвуют в расчёте; их отсутствие убирает соответствующий комментарий. Молча исключать уже допустимый актив из-за сбоя котировки запрещено.')

page('10 Источники и сбор данных')
table(['Источник','Контракт адаптера','Контроль'],[
('CBU FX API [1]','Запрос конкретной валюты и даты. Нормализовать Rate / Nominal. Сохранить дату действия курса.','Не использовать endpoint без даты как доказательство текущего курса'),
('CBU Gold [2]','Парсер опубликованной таблицы цен продажи и двух видов выкупа по весу слитка.','Не заявлять наличие отдельного публичного API без проверки'),
('Банки','Для demo — ручной каталог с проверкой. Production — только согласованные страницы банков и адаптеры.','У каждого продукта source_url, условия, checked_by, verified_at'),
('Допущенный криптопровайдер','Read only API или согласованный источник локального ask, шагов, минимумов и тарифов.','Нет подтверждённого API или прав доступа — crypto_enabled=false'),
('Инфляция','Официальная публикация статистики или ЦБ из утверждённого каталога источников.','Период показателя и дата публикации; не выдавать годовой показатель как прогноз'),
('НАПП [3]','Юридическое лицо, вид и номер лицензии, статус, домен, verified_at.','Неизвестный или недействующий статус отключает провайдера')],[1.4,3.0,2.6])
p('Все внешние источники подключаются через интерфейс fetch → parse → validate → persist. HTTP timeout 5 секунд, максимум 3 попытки на цикл с exponential backoff и jitter в общем лимите 30 секунд. Ошибка разбора не удаляет старые записи. После пяти подряд неудачных циклов источник переводится в degraded и отправляется сигнал эксплуатации.')
h('Проверки качества')
p('Отклонять нулевые и отрицательные цены, пустую валюту, неверные единицы, дубли ключей продукта и обратный порядок времени. Изменение цены более чем на 20% относительно предыдущей записи направлять в карантин для проверки вторым официальным источником или оператором. Это контроль качества, который не означает невозможность реального движения рынка.')
p('Сырые ответы хранить с SHA256 и parser_version. Нормализованные записи имеют уникальность source_id + instrument_id + kind + source_version. После успешной транзакции публиковать snapshot и кэш. Refresh lock исключает одновременный сбор одного источника несколькими воркерами.')
h('Границы доступности')
p('Библиотеки Yahoo Finance, Bloomberg, CoinGecko и Binance не являются обязательными зависимостями MVP. Лицензия данных, право парсинга, лимит запросов и доступность из среды размещения проверяются до подключения каждого источника. Фоновый сбор не создаёт заявки, кошельки или счета.')

page('11 Входной API и создание задания')
table(['Метод и путь с префиксом /api/v1','Результат'],[
('POST /auth/telegram','200 — сессия; raw init_data, auth_exchange_id UUID'),
('POST /portfolio/generate','202 — job; обязательный Idempotency-Key'),
('GET /portfolio/jobs/{job_id}','200 — состояние и готовый результат либо null'),
('POST /portfolio/jobs/{job_id}/clarify','202 — продолжение; expected_input_version, полный подтверждённый профиль'),
('POST /portfolio/jobs/{job_id}/cancel','200 — текущее состояние; повтор безопасен'),
('GET /portfolio/jobs/{job_id}/events','200 text/event-stream; авторизация и cursor'),
('GET /market/status','200 — времена и качество наборов, без персональных данных')],[3.5,3.5])
p('Все запросы, кроме auth, требуют Authorization: Bearer. JSON UTF-8, Content-Type application/json, максимальный body 16 KiB. UUID проверяются до обращения к БД. Неизвестные поля и неподдерживаемая schema_version отклоняются. OpenAPI является частью поставки.')
h('Пример GenerateRequestV1')
code('''{
  "schema_version": "1.0",
  "locale": "ru",
  "consent_version": "terms_v1",
  "privacy_version": "privacy_v1",
  "parent_job_id": null,
  "input": {
    "kind": "form", "capital_tiyin": "10000000000",
    "horizon_months": 12, "selected_risk": "moderate",
    "loss_capacity": 2, "loss_behavior": 1, "experience": 2,
    "max_loss_bps": 2000, "crypto_opt_in": true,
    "resident_uz": true, "adult_confirmed": true,
    "reserve_confirmed": true, "risk_acknowledged": true
  }
}''')
p('input — discriminated union. kind=text содержит только text длиной 1–2000 символов; остальные поля пользователь подтверждает через clarify. locale = ru | uz-Latn. parent_job_id = UUID|null и должен принадлежать текущему пользователю. Профиль из формы считается подтверждённым только при всех обязательных подтверждениях.')
code('''{
  "schema_version": "1.0", "job_id": "<uuid>",
  "status": "queued", "input_version": 1,
  "status_url": "/api/v1/portfolio/jobs/<uuid>",
  "events_url": "/api/v1/portfolio/jobs/<uuid>/events"
}''')

page('12 Выходной контракт и ошибки')
p('GET job всегда возвращает schema_version, job_id, status, stage, input_version, created_at, updated_at, clarification, result, error и warnings. clarification, result и error равны null вне соответствующего состояния. Статусы: queued, running, needs_input, completed, completed_with_warnings, failed, cancelled, expired. result появляется только в двух completed статусах.')
table(['PortfolioResultV1','Смысл и ограничение'],[
('allocation_id и snapshot_id','UUID; ссылки на неизменяемые расчёт и рынок'),
('policy_version и quant_version','Точные версии правил и алгоритма'),
('capital_tiyin и total_fee_tiyin','MoneyMinor string; сумма бюджета и явных комиссий'),
('cash_tiyin','MoneyMinor string; весь денежный остаток'),
('target_weights_bps','Все 4 asset_id; int 0–10000; сумма 10000 после ограничений'),
('positions[]','asset_id, product_id, quantity decimal string, quantity_unit, cost_tiyin, fee_tiyin, net_cost_share_bps; только ненулевые покупки'),
('cash_net_cost_share_bps','int; с долями positions даёт 10000'),
('gold_liquidation_tiyin','MoneyMinor string|null; оценка по известному bid с состоянием упаковки'),
('narrative_blocks[]','block_code и безопасный text; не источник финансовых полей'),
('warnings[]','code, message_code, params; тексты переводятся по словарю'),
('sources[]','source_id, source_url, observed_at|null, effective_date|null, fetched_at'),
('disclaimer','version, locale, статический text; обязательный ненулевой объект'),
('calculation_hash','SHA256 канонических входов и финансового результата')],[2.05,4.95])
code('''{
  "schema_version": "1.0",
  "code": "MARKET_DATA_UNAVAILABLE",
  "message": "Нет допустимой по свежести котировки золота",
  "field_errors": [], "retryable": true, "trace_id": "<uuid>"
}''')
p('HTTP до принятия job: 401 AUTH_INVALID, 404 NOT_FOUND для чужого и отсутствующего объекта, 409 IDEMPOTENCY_CONFLICT или INPUT_VERSION_CONFLICT, 413 BODY_TOO_LARGE, 422 VALIDATION_ERROR, 429 RATE_LIMITED с Retry-After, 503 STORAGE_UNAVAILABLE. После 202 ошибки выполнения находятся в error ресурса job; GET остаётся HTTP 200. Не возвращать stack trace, секреты или ответ провайдера в message.')

page('13 События и гарантии выполнения')
h('Идемпотентность')
p('Уникальность записи: user_id + endpoint + Idempotency-Key. Ключ — UUID, срок хранения 24 часа. Payload приводится к каноническому JSON и хешируется с версиями согласий и всеми значимыми параметрами. Тот же ключ и hash возвращают прежний job_id. Тот же ключ с другим hash возвращает 409. Новая попытка с новым ключом создаёт новый job и не переписывает результат прежнего.')
p('После авторизации сначала проверяются idempotency и hash; успешный повтор возвращается до списания квоты и active job limit. Job, ключ и outbox создаются одной транзакцией PostgreSQL. Worker берёт lease на 60 секунд с обновлением каждые 10 секунд. UNIQUE(job_id) и fencing_version исключают вторую публикацию. Dispatcher повторяет доставку outbox до подтверждения очереди.')
h('SSE')
p('Клиент использует streaming fetch с Authorization. Токен в URL запрещён. События сохраняются в PostgreSQL с UNIQUE(job_id, seq), а Redis доставляет уведомления. Порядок задаёт seq внутри job. Клиент повторно соединяется с Last-Event-ID и игнорирует уже применённые seq. Отключение клиента не отменяет расчёт.')
code('''id: 4
event: stage
data: {"schema_version":"1.0","stage":"quant_allocation",
       "message_code":"CALCULATING"}
''')
p('В реальном потоке data каждой записи отправляется одной строкой JSON; перенос выше иллюстративный. Разрешённые event: stage, warning, needs_input, completed, failed, cancelled, expired. Heartbeat — SSE comment каждые 15 секунд без seq. Финальное событие содержит job_id и result_url, но не сырой текст LLM. Nginx buffering отключается, Cache-Control: no-store.')
p('События хранятся 24 часа. Истёкший cursor получает HTTP 410 EVENTS_EXPIRED до открытия потока; клиент читает GET job. Максимальная длительность одного соединения 10 минут или до истечения сессии; затем повторная авторизация и reconnect. Проверка владельца обязательна на каждом соединении.')
h('Восстановление и отмена')
p('Checkpoint сохраняется после узла. При восстановлении worker сначала ищет предметные allocation/report по job_id. Итог, статус и финальное событие фиксируются одной транзакцией только при актуальном fencing_version, cancel_requested=false и now < deadline_at. Независимый sweeper каждые 2 секунды завершает просроченные job и отзывает lease. Статус failed появляется не позднее deadline + 2 секунды при доступной БД; публикация после deadline запрещена.')
p('needs_input освобождает worker и истекает через 24 часа. clarify требует expected_input_version и Idempotency-Key. cancel ставит cancel_requested, проверяемый между узлами и перед публикацией. После terminal статуса отмена возвращает текущий статус. Терминальные результаты неизменяемы.')

page('14 Схемы состояния и валидация')
p('На внешних границах использовать Pydantic с extra="forbid" и строгими типами. TypedDict внутри LangGraph описывает структуру для разработчика, но не выполняет runtime валидацию. Вход и выход каждого узла повторно проверяются схемой. Decimal создаётся из строки; не из float.')
code('''class PortfolioState(TypedDict):
    schema_version: str
    job_id: str
    input_version: int
    stage: str
    profile_ref: str | None
    snapshot_id: str | None
    allocation_id: str | None
    report_id: str | None
    clarification_ref: str | None
    errors: list[dict]
    warnings: list[dict]
    graph_version: str
    policy_version: str
    quant_version: str
    prompt_version: str
    model_config_version: str
    deadline_at: str
''')
table(['Контракт','Правило узла'],[
('ProfileDraftV1','Извлечённые поля nullable; evidence для каждого значения; missing_fields[] и conflicts[]. Нет автоназначенного риска.'),
('ConfirmedProfileV1','Полный профиль раздела 5, risk_acknowledged=true и версия согласий; effective_risk вычисляет код.'),
('MarketSnapshotV1','Поле snapshot_id закрепляется один раз; все котировки с единицами, временем и качеством.'),
('ValidatedAllocationV1','cost + cash + fees = capital; неотрицательность, целые лоты, hard caps и eligibility проверены.'),
('NarrativeV1','Только block_code, asset_id и explanation text; нет сумм, ставок и долей.'),
('ReportV1','Композиция ValidatedAllocation, принятого Narrative либо шаблона, источников и статического disclaimer.')],[1.65,5.35])
h('Идентичность и совместимость')
p('job_id обозначает один расчёт, thread_id — внутренний UUID исполнения, parent_job_id — предыдущий сценарий. thread_id клиент не задаёт. Соответствие thread_id → job_id → user_id хранит приложение. LangGraph PostgresSaver не выполняет проверку владельца.')
p('Версии закрепляются при создании job, snapshot — при первом чтении рынка. Изменение схемы требует миграции и теста возобновления. Неподдерживаемый старый checkpoint завершается RESTART_REQUIRED; продолжение с новыми правилами без нового job запрещено. Детерминированность относится к финансовому результату при одинаковых входах и версиях, а не к свободному тексту модели.')

page('15 Модель данных и сроки хранения')
table(['Таблица','Ключи и содержимое'],[
('users и sessions','UUID user; UNIQUE(telegram_user_id); hash токена, expiry, revoked_at'),
('consents','user_id, terms_version, privacy_version, risk_ack_version, accepted_at'),
('profiles','profile_id, user_id, profile_version, confirmed fields, effective_risk, policy_version'),
('portfolio_jobs','job_id, user_id, parent_job_id, thread_id, status, input_version, deadline, lease и версии'),
('job_inputs и clarifications','job_id + input_version; черновик и подтверждение отдельно; шифрование'),
('idempotency и outbox','UNIQUE(user_id, endpoint, key); job_id; canonical_hash; состояние доставки'),
('market_quotes','Временной ряд quote_id + recorded_at; source, product, kind, value, units, times, hash'),
('market_snapshots','snapshot_id и неизменяемая связь с quote_id, каталогом и реестром'),
('products и provider_registry','Версионированные условия, комиссии, availability, лицензия и verified_at'),
('allocations и reports','UNIQUE(job_id); финансовый JSON, narrative, disclaimer_version, hash'),
('job_events','UNIQUE(job_id, seq); тип, payload, created_at'),
('audit_events','actor, action, object_id, timestamp, trace_id, versions; без полного текста и секретов')],[1.65,5.35])
p('Пользовательские и транзакционные таблицы остаются обычными PostgreSQL таблицами. Timescale hypertable применяется только к котировкам; уникальные индексы включают колонку временного партиционирования. При расчёте не выполнять JOIN всей истории: snapshot напрямую ссылается на небольшой набор записей.')
h('Проектные сроки хранения')
p('Raw пользовательский текст — до 24 часов после терминального состояния; profile draft — до 24 часов; сессия — 15 минут; SSE и idempotency — 24 часа; raw рыночные ответы — 30 дней; нормализованные котировки — 3 года; профили, расчёты и аудит — 12 месяцев. Сроки являются конфигурацией retention_v1 и подлежат правовому согласованию до production.')
p('Удаление аккаунта скрывает историю сразу и удаляет персональные связи в течение 30 дней, если нет утверждённого законного основания удержания. Удалять или обезличивать checkpoint, предметные строки и файлы совместно. Резервные копии хранятся 30 дней; после восстановления повторно применять журнал удалений. Удаление части данных не должно оставлять доступ через parent_job_id.')
h('Миграции')
p('Alembic управляет предметной схемой. setup для PostgresSaver выполняется на этапе развёртывания с отдельными правами. Изменения каталога и политики создают новые версии. Редактирование принятого snapshot или уже опубликованного результата запрещено.')

page('16 Авторизация и безопасность')
h('Проверка Telegram initData')
p('Принять raw init_data, однократно разобрать query string, отклонить дубли ключей и некорректное кодирование. До проверки HMAC запрещено доверять user.id и initDataUnsafe. Для HMAC исключить hash, остальные полученные поля отсортировать и соединить key=value через LF без повторной сериализации user JSON. Применить официальный HMAC алгоритм Telegram [5].')
code('''secret = HMAC_SHA256(key=b"WebAppData", message=bot_token)
expected = HMAC_SHA256(key=secret, message=data_check_string)
valid = constant_time_compare(expected.hex(), received_hash)
''')
p('auth_date не старше 300 секунд, будущее смещение не более 30 секунд. user.id обязателен. Digest initData хранить до auth_date + 300 секунд. Повтор того же auth_exchange_id и digest возвращает сессию в защищённом 30 секундном кэше. Другой digest с тем же ID даёт 409 AUTH_EXCHANGE_CONFLICT, тот же digest с другим ID — 409 AUTH_REPLAY. После окна повтора требуется заново открыть Mini App.')
p('Сессия — случайный непрозрачный bearer token не менее 256 бит, срок 15 минут. Backend хранит только hash, а исходный токен кратковременно шифруется для повтора auth exchange. Клиент хранит токен в памяти. Токен бота и ключи провайдеров хранятся в менеджере секретов и доступны только нужным контейнерам.')
h('Контроль доступа и нагрузки')
table(['Контроль','Политика MVP'],[
('Ownership','Каждый read, clarify, cancel, SSE и parent_job_id фильтруется по user_id сессии'),
('Generate','5 запросов в минуту, 30 в сутки на пользователя; не более 2 активных job'),
('Clarify и auth','10 clarify в минуту на пользователя; auth 20 запросов в минуту на IP'),
('SSE','Не более 3 одновременных соединений на пользователя'),
('Redis outage','Ограниченный DB limiter; при невозможности безопасно ограничить — 503'),
('Сеть','TLS, точный CORS allowlist; БД и Redis без публичных портов'),
('Журналы','Не писать токены, initData, имя, телефон, raw prompt и полный профиль')],[1.5,5.5])
p('Все ID считаются недоверенными. Чужой объект возвращает 404. Параметризованные SQL обязательны. Очищать Markdown, отключить raw HTML, разрешать ссылки только из каталога источников. LLM не имеет сетевого инструмента, Python interpreter или доступа к пользовательским сессиям.')

page('17 Правовые требования к запуску')
p('До публичного запуска владелец продукта совместно с юристом определяет правовой режим фактических функций сервиса, применимые уведомления, лицензии и ограничения для каждого инструмента. Персонализация по капиталу и риску оценивается по содержанию услуги. Название «симулятор» и наличие дисклеймера сами по себе не снимают регуляторные обязанности.')
h('Провайдеры и криптовалюта')
p('В проверенном 22 октября 2026 года реестре НАПП лицензия KOBEA GROUP CE №0003 аннулирована 25 мая 2026 года приказом №45 [3]. Привязка сценария к UzNEX из исходного примера не используется. Система хранит допущенное юридическое лицо, номер и вид лицензии, домен и дату проверки. Статус unknown, suspended или revoked блокирует crypto; действующая лицензия сама по себе не подтверждает доступность API.')
p('MVP ограничивает аудиторию совершеннолетними резидентами РУз как продуктовым условием. Допуск клиента к конкретной операции, договор и KYC выполняет провайдер. Сервис не собирает паспорт и биометрию и не хранит приватные ключи кошельков. Правила для нерезидентов не обобщаются из правил MVP.')
h('Персональные данные и внешняя LLM')
p('Статья 27¹ закона о персональных данных была изменена в марте 2026 года. Действующая редакция различает обязательное хранение отдельных категорий в РУз и условия зарубежной обработки других данных [4]. Для сервиса нужны классификация данных, проверка оснований, стран и договоров обработки; старое обобщение о локализации всех данных не применяется автоматически.')
p('Базовое размещение пользовательской БД и резервных копий — инфраструктура в Узбекистане. До письменного допуска внешней обработки использовать форму и локальный шаблон отчёта. Внешняя LLM выключена. После допуска передавать только минимальные поля без Telegram ID, имени, телефона, session token и исходного текста с идентификаторами. Псевдонимизация не считается автоматическим освобождением от требований закона.')
h('Статический дисклеймер')
p('Обязательный текст версии disclaimer_v1 для тестового MVP: «Материал носит информационный характер и не является индивидуальной инвестиционной рекомендацией. Стоимость активов может снижаться. Доходность в прошлом не гарантирует доходности в будущем. Расчёт основан на данных и условиях на указанную дату; комиссии, наличие и условия необходимо проверить у поставщика. Операции с криптоактивами требуют соблюдения законодательства Республики Узбекистан и условий допущенного провайдера».')
p('Окончательные тексты ru и uz-Latn утверждаются до production. Backend всегда добавляет дисклеймер; frontend показывает его до действий по сценарию. Запрещены фразы «безрисковый депозит», «гарантированная доходность» и безусловное обещание защиты от девальвации. Условия гарантий вкладов показываются только по отдельным проверенным правилам.')

page('18 Контроль LLM и формирование отчёта')
p('В форме LLM не вызывается для профилирования. Для текста используется structured output ProfileDraftV1, а финансовое подтверждение выполняет пользователь. Модель не может достроить возраст, инвестиционный опыт, согласие или риск по стилю сообщения. Обработка исходного текста внешней моделью допускается только после выполнения условий раздела 17.')
h('Схема пояснения')
code('''{
  "schema_version": "1.0",
  "blocks": [
    {"block_code": "ROLE_OF_ASSET", "asset_id": "gold_cbu",
     "text": "Возможны колебания цены и потери на спреде."}
  ]
}''')
p('Разрешённые block_code: SUMMARY, ROLE_OF_ASSET, RISKS, NEXT_STEPS. asset_id обязателен для ROLE_OF_ASSET, иначе null. Максимум 8 блоков, text до 600 символов, весь narrative до 3000 символов. Модель получает только разрешённые факты и коды активов. Суммы и проценты из её ответа не используются.')
h('Проверка результата модели')
p('Отклонять неизвестные asset_id, цифры, знак процента, числительные с финансовыми величинами, обещание дохода, неразрешённые названия продуктов, ссылки и требования совершить сделку. После схемы применяется семантический фильтр и регрессионный набор запрещённых утверждений на обоих языках. При сомнении применяется проверенный статический шаблон. Сам по себе промпт «не меняй цифры» защитой не считается.')
p('Числовые таблицы, шаги открытия продукта с проверенной ссылкой, предупреждения и дисклеймер формирует код. В пользовательский интерфейс не стримится непроверенный текст модели. Запросы пользователя и тексты внешних страниц считаются данными и не могут изменить инструкции, источники, лимиты или инструменты.')
p('Прогноз доходности портфеля в MVP не выдаётся. Для справочного сравнения за один и тот же период код использует real_return = (1 + nominal_return) / (1 + inflation) - 1. Условные 22% и 9% дают около 11,93%, а не 13%. Денежные потоки депозита считаются только по известным условиям договора.')
h('Модель и эксплуатационный бюджет')
p('Модель подключается через адаптер и конфигурацию с provider, точным model_id, region, timeout и версией prompt. Упомянутый в концепте gemini-1.5-pro отключён провайдером 29 сентября 2025 года [8]. Использовать его в поставке запрещено. При развёртывании проверять доступность выбранной поддерживаемой модели и structured output; смена модели проходит тот же набор evals.')
p('На LLM вызов 10 секунд, максимум одна повторная попытка, только если остаётся общий бюджет. Некорректная схема допускает одну исправляющую попытку в том же лимите. Не более двух обращений для одного узла. Отказ профилировщика ведёт к форме; отказ пояснения — к шаблону. Температура и ограничения токенов фиксируются в model_config_version.')

page('19 Развёртывание и эксплуатационные показатели')
p('Референсный стенд: API 2 vCPU и 2 GiB RAM; worker 4 vCPU и 4 GiB RAM, concurrency=4; PostgreSQL с TimescaleDB 4 vCPU и 8 GiB RAM на SSD; Redis 1 vCPU и 1 GiB RAM. Это база для нагрузочных испытаний, не обещание внешнего SLA. Все компоненты в одной сети; лимиты внешних API и LLM фиксируются в протоколе испытаний.')
table(['Показатель','Цель и способ измерения'],[
('Приём generate','p95 ≤ 500 мс до 202; без LLM в HTTP handler'),
('Первый SSE event','p95 ≤ 1 секунды после подключения; сохранённое событие доступно сразу'),
('Готовый результат формы','p95 ≤ 30 секунд при 20 параллельных job; deadline 45 секунд с ожиданием очереди'),
('Quant Core','p95 ≤ 500 мс, hard timeout 2 секунды на заданных пределах бюджета'),
('Доступность API','Целевая 99,5% в месяц; мониторинг отдельно от доступности источников'),
('RPO и RTO','RPO ≤ 15 минут, RTO ≤ 2 часов при настроенном PITR и проверенном восстановлении'),
('Deadline и fallback','Перед истечением срока вернуть шаблон, если расчёт уже валиден; иначе DEADLINE_EXCEEDED')],[1.8,5.2])
p('Для needs_input время ожидания пользователя не включается в 45 секунд. После принятого clarify начинается новый интервал 45 секунд. Повторные попытки узлов и ожидание очереди входят в этот интервал. Redis outage допускает задержку очереди до восстановления; пропущенный deadline завершается явно, а не теряется.')
h('Наблюдаемость')
p('Структурированные журналы содержат trace_id, job_id, stage, latency, error_code и версии без содержимого профиля. Метрики: очередь и её возраст, outbox backlog, ошибки парсеров, возраст snapshot, частота static fallback, ошибки HMAC и ownership, invariant violations, длительность Quant и стоимость LLM. INVARIANT_VIOLATION вызывает немедленный alert и блокирует публикацию.')
p('Предупреждение эксплуатации при outbox старше 30 секунд, p95 job выше 30 секунд в течение 5 минут, недоступном источнике выше жёсткого возраста или fallback выше 10% за 15 минут при не менее 20 job. Health API: /health/live проверяет процесс, /health/ready — доступность БД и применённые миграции. Состояние LLM не делает API неготовым, если доступен шаблон.')
h('Резервирование и обновление')
p('Ежедневный backup, WAL/PITR и ежемесячная проверка восстановления на изолированном стенде. Релиз проходит миграцию, smoke и проверку возобновления старого job. Откат приложения допускается только при совместимой схеме. Пользовательские данные не отправляются во внешнюю систему трассировки без отдельного согласования.')

page('20 Проверка расчётов и качества данных')
p('Критерии ниже проверяются автоматически на замороженных fixtures. Тесты используют один и тот же snapshot и версии политики. Числовое совпадение требуется до тийина, а не с допуском float. Для инвариантов применяются параметризованные и генеративные проверки.')
table(['ID','Вход или сбой','Ожидаемый результат'],[
('Q01','Фикстура F100 из раздела 8','Точные суммы, 0.01485148 BTC; cash 2 000 005,20; комиссии 148 514,80 UZS'),
('Q02','Капитал с остатком при делении на 10000','Риск бюджеты округлены вниз; дробный остаток в cash; сумма равна капиталу'),
('Q03','Переставить входные записи каталога','Финансовый результат и calculation_hash не меняются'),
('Q04','Золото ниже минимального слитка','0 слитков, весь бюджет класса в cash, причина BELOW_MIN_LOT'),
('Q05','Следующий шаг BTC вместе с комиссией выше бюджета','Берётся предыдущий шаг; ни один тийин не заимствован у другого класса'),
('Q06','BTC ниже min_notional','Количество и комиссия 0, бюджет в cash'),
('Q07','Conservative и текст «добавь крипты»','BTC 0; текст не переопределяет policy'),
('Q08','Границы горизонта 5/6, 11/12, 23/24','Применяются заданные caps и effective_risk; запрещённый актив не появляется'),
('Q09','Профиль с неполными или конфликтными ответами','needs_input, Quant не запускается'),
('Q10','Нарушить сумму, лот или hard cap вручную','INVARIANT_VIOLATION, отчёт не публикуется'),
('D01','Котировка между мягкой и жёсткой границей','Разрешён старый срез с STALE_DATA_USED и реальной датой'),
('D02','Котировка старше жёсткой границы','MARKET_DATA_UNAVAILABLE, без подстановки demo'),
('D03','Парсер вернул пустую таблицу или цену 0','Новый snapshot не опубликован; старый сохранён'),
('D04','FX с Nominal=100','В расчёте Rate / 100; effective_date не равна fetched_at автоматически'),
('D05','Лицензия revoked или проверка старше 24 часов','BTC исключён на eligibility, причина отражена в result'),
('D06','Redis удалён после публикации snapshot','Срез читается из PostgreSQL, значения не меняются')],[.55,2.9,3.55])
p('Проверить max_loss_bps на 999/1000, 1999/2000, 2499/2500; experience<2 и selected_risk ниже анкеты. Проверить gold ровно 25% C при комиссии BTC: gross cap соблюдён, net share не используется как cap. На 10000 случайных входов проверить баланс и сумму отображаемых долей 10000. Эти проверки не заменяют F100.')

page('21 Приёмка API безопасности и интерфейса')
table(['ID','Проверка','Критерий прохождения'],[
('A01','20 одновременных generate с одним ключом','Один job, одна запись результата, одинаковый job_id во всех ответах'),
('A02','Тот же ключ с другой суммой','409 IDEMPOTENCY_CONFLICT'),
('A03','Одновременные clarify одной версии','Одна принятая новая input_version; другой запрос получает 409 либо идемпотентный повтор'),
('A04','Worker аварийно остановлен после Quant','Возобновление с тем же snapshot; один опубликованный результат'),
('A05','Потеря Redis и повтор outbox','Job остаётся в БД; повторная доставка не создаёт вторую публикацию'),
('A06','SSE разорван после seq=4','Reconnect возвращает seq>4 по порядку; terminal status совпадает с GET'),
('S01','Подменить initData, user JSON или auth_date','401 до выдачи сессии; тесты официального HMAC алгоритма'),
('S02','Чужие job, parent_job_id, SSE и clarify','404, никакие поля другого пользователя не возвращаются'),
('S03','Токены и raw текст в запросах','Секреты отсутствуют в логах, checkpoint, URL и внешней трассировке'),
('S04','Превысить rate limit и active job limit','429 и Retry-After; валидный повтор с ключом не создаёт новый job'),
('L01','LLM вернула «гарантировано» или цифры','Текст отвергнут; шаблон, completed_with_warnings'),
('L02','LLM недоступна или схема неверна','Форма остаётся работоспособной, проверенный результат доступен'),
('U01','360 px, Android и iOS Telegram','Форма, уточнение, таблица и disclaimer доступны без перекрытия'),
('U02','Переключить ru на uz-Latn','Суммы, доли, даты и обязательный текст соответствуют исходному контракту'),
('O01','Нагрузочный прогон','Показатели раздела 19 подтверждены протоколом с квотами и версией модели'),
('O02','Восстановление и удаление аккаунта','RPO/RTO измерены; после restore удалённый пользователь не получает историю')],[.55,2.9,3.55])
p('Приёмка MVP требует прохождения всех Q, D, A, S, L, U и O проверок. Для демонстрации crypto допускается только тестовый каталог, явно обозначенный в интерфейсе. Подтверждение испытаний не означает разрешения публичной финансовой услуги; условия запуска перечислены в разделе 22.')

page('22 Поставка и условия публичного запуска')
h('Структура репозитория')
code('''uzwealth/
  app/api/v1/          # auth, portfolio, events, market
  app/schemas/        # versioned Pydantic contracts
  app/graph/          # state, nodes, workflow
  app/quant/          # policy, budgets, lots, fees, invariants
  app/market/         # adapters, catalog, freshness, snapshots
  app/security/       # Telegram validation, sessions, ownership
  app/workers/        # ingestion, outbox, jobs, retention
  app/templates/      # ru, uz-Latn, static legal text
  migrations/         # application and checkpoint migrations
  frontend/           # Mini App React TypeScript
  tests/fixtures/     # F100 and source parser samples
  tests/              # quant, API, failure, security, E2E
  deploy/             # compose, proxy, health, backup runbook
  docs/               # OpenAPI, source registry, ADR, operations
''')
p('Поставка включает исходники, lock файл зависимостей, Docker образы и Compose для стенда, Alembic миграции, OpenAPI JSON, fixtures, отчёт тестов, инструкции запуска и восстановления, реестр источников и конфигурацию политик. Точные совместимые версии Python, LangGraph, PostgresSaver, FastAPI, Pydantic, PostgreSQL и TimescaleDB фиксируются в lock и image digest после интеграционного прогона. Код из концепта не считается готовым production компонентом.')
h('Этапы разработки')
p('Этап 1: схемы, чистый Quant Core и F100. Этап 2: каталог, ingestion, snapshot и контроль свежести. Этап 3: API, авторизация, outbox, checkpoint и SSE. Этап 4: Mini App, языки, шаблоны и опциональный LLM адаптер. Этап 5: сбои, нагрузка, безопасность и восстановление. Каждый этап закрывается соответствующими проверками разделов 20–21; календарные сроки определяются командой после оценки работ.')
table(['Условие production','Ответственный и подтверждение'],[
('Методология риска','Владелец риска утверждает матрицу, анкету, пороги и тексты без обещания просадки'),
('Правовая квалификация','Юрист и владелец продукта определяют режим сервиса, согласия, disclaimer и права на данные'),
('Персональные данные','Владелец данных утверждает размещение, внешнюю LLM, страны, основания и retention'),
('Провайдеры и источники','Операции подтверждают API, тарифы, доступ и действующий статус; crypto включается отдельно'),
('Эксплуатация','Техлид подтверждает SLO, backup, мониторинг, квоты и ответственных за инциденты')],[2.0,5.0])

page('23 Источники и контроль изменений')
p('Первичные источники проверены 22 октября 2026 года. Ссылки задают основания и интерфейсы; рыночные значения из страниц не фиксируются как постоянные параметры продукта. Перед production повторно проверить реестр, редакции актов, условия источников и поддержку библиотек.')
sources=[
('[1] ЦБ РУз Форматы валютных курсов','https://cbu.uz/ru/arkhiv-kursov-valyut/veb-masteram/','Параметры даты и валюты, Rate, Nominal и Date. Требования разделов 9–10.'),
('[2] ЦБ РУз Продажа и выкуп золотых слитков','https://cbu.uz/ru/banknotes-coins/gold-bars/prices/','Раздельные цены ask и выкупа, веса и состояние упаковки. Разделы 7–10.'),
('[3] НАПП Реестр провайдеров криптоактивов','https://napp.uz/ru/pages/erlsp','Статусы и история лицензий, включая KOBEA GROUP CE №0003. Разделы 10 и 17.'),
('[4] Закон о персональных данных','https://lex.uz/docs/4396428','Действующая редакция и ст 27¹; русский текст изменений отмечен LexUZ как неофициальный перевод.'),
('Официальный узбекский текст закона','https://lex.uz/docs/4396419','Использовать для правовой проверки вместе с актуальными изменениями. Разделы 15–17.'),
('[5] Telegram Mini Apps','https://core.telegram.org/bots/webapps#validating-data-received-via-the-mini-app','Серверная проверка initData и auth_date. Раздел 16.'),
('[6] LangGraph Persistence','https://docs.langchain.com/oss/python/langgraph/persistence','compile с checkpointer и привязка к thread_id. Разделы 4 и 14.'),
('[7] LangGraph Checkpointers','https://docs.langchain.com/oss/python/langgraph/checkpointers','Хранение и восстановление состояния. Разделы 13–15.'),
('[8] Google Gemini API Release notes','https://ai.google.dev/gemini-api/docs/changelog','Отключение моделей Gemini 1.5 29 сентября 2025 года. Раздел 18.')]
for title,url,note in sources:
    par=p(title);par.paragraph_format.keep_with_next=True
    r=par.runs[0];r.bold=True;r.font.size=Pt(10)
    par=p(url,'Small');par.paragraph_format.keep_with_next=True
    p(note,'Small')
h('Порядок изменения ТЗ')
p('Изменение схем, правил риска, источника цены, тарифов или дисклеймера получает новую версию и запись ADR с причиной, владельцем и затронутыми тестами. Старые job остаются воспроизводимыми по закреплённым версиям. Изменение модели не может изменить финансовый расчёт. Переход к исполнению сделок или новым классам активов требует расширения границ продукта и отдельной приёмки.')

page('24 Контракты уточнения профиля')
p('Контракты этого приложения уточняют разделы 11–14. MoneyMinor — строка по regex ^(0|[1-9][0-9]*)$, для capital дополнительно действуют пределы раздела 5. DecimalString — неотрицательная десятичная строка без экспоненты; money и quantity не преобразуются через JavaScript Number. bool не принимается вместо int. Даты — RFC3339 UTC, даты действия — YYYY-MM-DD.')
h('ProfileDraftV1 и ClarificationV1')
p('ProfileDraftV1 содержит все поля формы из GenerateRequestV1, кроме kind. Каждое поле nullable; дополнительные поля запрещены. Неоднозначное или не найденное значение — null. evidence[] содержит field, exact_quote, start_offset и end_offset в символах исходного текста; отсутствие доказательства оставляет поле null. Внешняя LLM не получает текст до допуска его обработки.')
table(['Поле ClarificationV1','Тип и правило'],[
('clarification_id','UUID, создаётся backend для каждого уточнения'),
('reason_code','PROFILE_CONFIRMATION_REQUIRED | PROFILE_INCOMPLETE | PROFILE_CONFLICT | SUITABILITY_NOT_CONFIRMED | PARSER_UNAVAILABLE'),
('draft','ProfileDraftV1; все ключи присутствуют, пропуски null'),
('missing_fields','Массив имён полей формы без повторов'),
('field_errors','Массив объектов field, code, message_code; code = REQUIRED | CONFLICT | UNSUITABLE | OUT_OF_RANGE'),
('expected_input_version','int ≥ 1, текущая версия job'),
('confirmation_required','Всегда true; фронтенд показывает полную форму подтверждения')],[2.1,4.9])
h('ClarifyRequestV1')
code('''{
  "schema_version": "1.0",
  "clarification_id": "11111111-1111-4111-8111-111111111111",
  "expected_input_version": 1,
  "consent_version": "terms_v1", "privacy_version": "privacy_v1",
  "profile": {
    "capital_tiyin": "10000000000", "horizon_months": 12,
    "selected_risk": "moderate", "loss_capacity": 2,
    "loss_behavior": 1, "experience": 2, "max_loss_bps": 2000,
    "crypto_opt_in": true, "resident_uz": true,
    "adult_confirmed": true, "reserve_confirmed": true,
    "risk_acknowledged": true
  }
}''')
p('Backend валидирует профиль, согласия и владельца. Атомарно проверить status=needs_input, clarification_id и input_version, сохранить новую версию и outbox, увеличить input_version и перевести job в queued. Иная версия — 409 INPUT_VERSION_CONFLICT; иной статус — 409 JOB_STATE_CONFLICT.')
p('В запросе clarify действует Idempotency-Key. Точный успешный повтор возвращает прежний ответ до проверки изменившегося состояния job. Новый запрос не может менять running или terminal job. Несогласованные обязательные подтверждения возвращают needs_input с новой clarification_id и не запускают чтение рынка.')

page('25 Состояния задания и контроль сумм')
table(['Статус','Допустимое содержимое и переход'],[
('queued','clarification/result/error = null; переход в running, cancelled или failed'),
('running','Все три null; переход в needs_input, completed, completed_with_warnings, failed или cancelled'),
('needs_input','clarification обязателен; result/error=null; clarify → queued; cancel → cancelled; 24 часа → expired'),
('completed','result обязателен, warnings пуст; clarification/error=null; terminal'),
('completed_with_warnings','result обязателен, warnings непуст; clarification/error=null; terminal'),
('failed','error обязателен; result/clarification=null; terminal'),
('cancelled','error/result/clarification=null; terminal'),
('expired','error.code=INPUT_EXPIRED; result/clarification=null; terminal')],[1.7,5.3])
p('stage — enum queued | extract_profile | get_market_data | quant_allocation | generate_report | awaiting_input | done. Для queued stage=queued; needs_input=awaiting_input; terminal=done. running использует только четыре имени узлов. Каждое событие stage содержит текущий status. Событие completed обязательно содержит окончательный status и result_url, чтобы клиент различал предупреждения.')
h('Инварианты ValidatedAllocationV1')
code('''total_fee_tiyin = sum(position.fee_tiyin)
sum(position.cost_tiyin) + cash_tiyin + total_fee_tiyin = C
D = C - total_fee_tiyin
assert D > 0
# Для каждой позиции, для cash отдельный числитель:
base_bps[i] = numerator_tiyin[i] * 10000 // D
rem[i] = numerator_tiyin[i] * 10000 % D
# Недостающие bps: rem по убыванию, asset_id по возрастанию.
sum(net_cost_share_bps) + cash_net_cost_share_bps = 10000
''')
p('numerator равен cost_tiyin актива, для cash — cash_tiyin. Все cost, fee, cash и количества неотрицательны. Количества кратны шагам и stock limits. Hard cap каждого класса проверяется по cost + fee относительно C, до округления процентов. Отсутствие подходящего продукта отличается от отсутствия актуальной цены уже допущенного продукта: первое даёт cash, второе блокирует расчёт.')
p('Для депозитов и золота MVP допускает фиксированную известную комиссию в UZS за продукт или слиток. Применимые комиссии округляются вверх до тийина. Депозит и золото не могут расходовать бюджет соседнего класса. Недопущенные продукты и источники не становятся допустимыми из-за пожелания LLM.')
p('calculation_hash вычисляется из канонических бизнес полей подтверждённого профиля, hash содержимого snapshot, policy_version, quant_version и финансовых выходов. Исключить job_id, время подтверждения и текст LLM; массивы сортировать по asset_id и product_id. Повтор с теми же бизнес входами даёт тот же hash.')

doc.save(OUT)
print(json.dumps({'output':str(OUT),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)},ensure_ascii=False))
